"""内容导出引擎：支持 Markdown / TXT / DOCX / PDF 四种格式。"""
import io
import logging
import os
import re
from pathlib import Path
from typing import Optional

from bs4 import BeautifulSoup
from docx import Document
from fpdf import FPDF
from markdown import markdown

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

logger = logging.getLogger(__name__)


class Exporter:
    def __init__(self, project_root: Path):
        self.project_root = project_root
        self.pdf_font_path = self._detect_font_path()

    def _detect_font_path(self) -> Optional[str]:
        env_font = os.getenv("PDF_FONT_PATH")
        if env_font and Path(env_font).exists():
            return env_font
        candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
            "/System/Library/Fonts/Supplemental/ArialUnicode.ttf",
            "/System/Library/Fonts/STHeiti Light.ttc",
            "/Library/Fonts/Arial Unicode.ttf",
            "C:/Windows/Fonts/arialuni.ttf",
        ]
        for path in candidates:
            if Path(path).exists():
                logger.info(f"PDF 字体: {path}")
                return path
        logger.warning("未找到 Unicode 字体，PDF 可能不支持非 ASCII 字符")
        return None

    def markdown_to_plain(self, content: str) -> str:
        html = markdown(content or "")
        soup = BeautifulSoup(html, "html.parser")
        text = soup.get_text("\n")
        lines = [line.rstrip() for line in text.splitlines()]
        plain = "\n".join(lines).strip()
        return re.sub(r"\n{2,}", "\n", plain)

    def export_markdown(self, content: str) -> io.BytesIO:
        buf = io.BytesIO()
        buf.write((content or "").encode("utf-8"))
        buf.seek(0)
        return buf

    def export_text(self, content: str) -> io.BytesIO:
        plain = self.markdown_to_plain(content or "")
        buf = io.BytesIO()
        buf.write(plain.encode("utf-8"))
        buf.seek(0)
        return buf

    def export_docx(self, content: str) -> io.BytesIO:
        plain = self.markdown_to_plain(content or "")
        doc = Document()
        for line in (plain.splitlines() or [""]):
            doc.add_paragraph(line if line.strip() else "")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def export_pdf(self, content: str) -> io.BytesIO:
        plain = self.markdown_to_plain(content or "")
        if REPORTLAB_AVAILABLE:
            try:
                return self._pdf_reportlab(plain)
            except Exception as exc:
                logger.error(f"ReportLab 失败，回退 FPDF: {exc}")
        return self._pdf_fpdf(plain)

    def _pdf_reportlab(self, plain: str) -> io.BytesIO:
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        font_name = self._register_rl_font()
        x_margin, top, bottom = 40, A4[1] - 50, 40
        text_obj = c.beginText(x_margin, top)
        text_obj.setFont(font_name, 12)
        for line in (plain.splitlines() or [""]):
            if text_obj.getY() <= bottom:
                c.drawText(text_obj)
                c.showPage()
                text_obj = c.beginText(x_margin, top)
                text_obj.setFont(font_name, 12)
            text_obj.textLine(line)
        c.drawText(text_obj)
        c.save()
        buf.seek(0)
        return buf

    def _register_rl_font(self) -> str:
        font_name = "CustomRL"
        try:
            if self.pdf_font_path and Path(self.pdf_font_path).exists():
                pdfmetrics.registerFont(TTFont(font_name, self.pdf_font_path))
                return font_name
        except Exception:
            pass
        return "Helvetica"

    def _pdf_fpdf(self, plain: str) -> io.BytesIO:
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        if self.pdf_font_path and self.pdf_font_path.lower().endswith((".ttf", ".otf")):
            try:
                pdf.add_font("Custom", "", self.pdf_font_path, uni=True)
                pdf.set_font("Custom", size=12)
            except Exception:
                pdf.set_font("Helvetica", size=12)
        else:
            pdf.set_font("Helvetica", size=12)

        max_w = pdf.w - pdf.l_margin - pdf.r_margin
        for line in (plain.splitlines() or [""]):
            safe = line.replace("\t", "    ")
            self._fpdf_write_line(pdf, safe, max_w)
        pdf_bytes = pdf.output(dest="S")
        if isinstance(pdf_bytes, str):
            pdf_bytes = pdf_bytes.encode("latin1")
        buf = io.BytesIO(pdf_bytes)
        buf.seek(0)
        return buf

    def _fpdf_write_line(self, pdf: FPDF, text: str, max_w: float):
        buf = ""
        for ch in text:
            if ch == "\n":
                if buf:
                    pdf.cell(0, 8, buf, ln=1)
                    buf = ""
                pdf.cell(0, 8, " ", ln=1)
                continue
            if pdf.get_string_width(buf + ch) <= max_w:
                buf += ch
            else:
                if buf:
                    pdf.cell(0, 8, buf, ln=1)
                buf = ch
        if buf:
            pdf.cell(0, 8, buf, ln=1)
